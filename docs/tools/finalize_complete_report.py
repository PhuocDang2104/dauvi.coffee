from __future__ import annotations

import argparse
import os
import shutil
import tempfile
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from finalize_project_report import Writer, clear_between, find_paragraph


BODY_STYLES = {"Normal", "List Paragraph", "No Spacing"}
TABLE_HEADER_GRAY = "D9D9D9"
COMMAND_GRAY = "D9D9D9"
DEMO_URL = "https://dauvi-coffee.vercel.app/"
DIAGRAM_MEDIA = {
    "word/media/image5.png": "06-use-case-overview.png",
    "word/media/image6.png": "07-system-architecture-rag.png",
    "word/media/image7.png": "08-login-sequence.png",
    "word/media/image8.png": "09-order-sequence.png",
    "word/media/image9.png": "10-chatbot-sequence.png",
    "word/media/image10.png": "11-admin-sequence-target.png",
    "word/media/image11.png": "13-current-erd-pgvector.png",
    "word/media/image12.png": "12-langgraph-workflow.png",
}


def remove_paragraph(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def set_update_fields(document: Document) -> None:
    settings = document.settings._element
    current = settings.find(qn("w:updateFields"))
    if current is None:
        current = OxmlElement("w:updateFields")
        settings.append(current)
    current.set(qn("w:val"), "true")


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def fix_preface_layout(document: Document) -> None:
    heading = find_paragraph(document, "LỜI MỞ ĐẦU")
    body_count = 0
    current = heading._p.getnext()
    while current is not None and body_count < 4:
        if current.tag == qn("w:p"):
            paragraph = next(p for p in document.paragraphs if p._p is current)
            if paragraph.style.name == "CAP 1":
                break
            if paragraph.text.strip():
                clean_text = paragraph.text.lstrip(" \t\u00a0")
                paragraph.clear()
                paragraph.add_run(clean_text)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                paragraph.paragraph_format.left_indent = Cm(0)
                paragraph.paragraph_format.right_indent = Cm(0)
                paragraph.paragraph_format.first_line_indent = Cm(1.27)
                paragraph.paragraph_format.line_spacing = 1.5
                paragraph.paragraph_format.space_after = Pt(0)
                body_count += 1
        current = current.getnext()
    if body_count < 2:
        raise RuntimeError("Không tìm thấy đủ hai đoạn đầu của LỜI MỞ ĐẦU")


def standardize_table_colors(document: Document) -> None:
    for table in document.tables:
        fills = []
        for row in table.rows:
            for cell in row.cells:
                shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
                fills.append(shading.get(qn("w:fill")) if shading is not None else None)

        is_command_block = (
            len(table.rows) == 1
            and len(table.columns) == 1
            and any(fill in {"F4F1EA", COMMAND_GRAY} for fill in fills)
        )
        has_colored_header = bool(fills) and fills[0] in {"173629", TABLE_HEADER_GRAY}

        if is_command_block:
            cell = table.cell(0, 0)
            shade_cell(cell, COMMAND_GRAY)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)

        if has_colored_header:
            for cell in table.rows[0].cells:
                shade_cell(cell, TABLE_HEADER_GRAY)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0, 0, 0)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "Times New Roman")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "26")
    properties.extend([fonts, color, underline, size])
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def ensure_demo_link(document: Document) -> None:
    existing_urls = {
        relationship.target_ref
        for relationship in document.part.rels.values()
        if relationship.reltype.endswith("/hyperlink")
    }
    if DEMO_URL in existing_urls:
        return
    deployment_heading = find_paragraph(document, "4.6.1. Frontend trên Vercel")
    paragraph_element = deployment_heading._p.getnext()
    while paragraph_element is not None and paragraph_element.tag != qn("w:p"):
        paragraph_element = paragraph_element.getnext()
    if paragraph_element is None:
        raise RuntimeError("Không tìm thấy đoạn triển khai Vercel để chèn link demo")
    paragraph = next(p for p in document.paragraphs if p._p is paragraph_element)
    paragraph.add_run(" Website demo: ")
    add_hyperlink(paragraph, DEMO_URL, DEMO_URL)


def fix_template_placeholders(document: Document) -> None:
    def replace_text(paragraph, text: str) -> None:
        run_properties = (
            deepcopy(paragraph.runs[0]._r.rPr)
            if paragraph.runs and paragraph.runs[0]._r.rPr is not None
            else None
        )
        paragraph.clear()
        run = paragraph.add_run(text)
        if run_properties is not None:
            current = run._r.rPr
            if current is not None:
                run._r.remove(current)
            run._r.insert(0, run_properties)

    replacements = {
        "Thiết Kế và Xây dựng trang web: Bán cà phê": (
            "THIẾT KẾ VÀ XÂY DỰNG WEBSITE BÁN CÀ PHÊ TÍCH HỢP CHATBOT AI"
        ),
        "\t\tNguyễn Thị A": "\t\tLưu Phạm Vĩnh Tùng",
        "THIẾT KẾ VÀ XÂY DỰNG WEBSITE HỖ TRỢ HỌC TẬP": (
            "THIẾT KẾ VÀ XÂY DỰNG WEBSITE BÁN CÀ PHÊ TÍCH HỢP CHATBOT AI"
        ),
    }
    for paragraph in document.paragraphs:
        if paragraph.text in replacements:
            replace_text(paragraph, replacements[paragraph.text])
        elif "MSSV: 2200001234" in paragraph.text:
            replace_text(
                paragraph,
                "\tHọ và tên: LƯU PHẠM VĨNH TÙNG\tMSSV: 2311557532\t",
            )
        elif "Lớp: 22DKTPM" in paragraph.text:
            replace_text(
                paragraph,
                "\tChuyên ngành: Kỹ thuật công nghệ thông tin\tLớp: 23DTH1A\n"
                "\tEmail: …………………………………\tSĐT: 0396100730\n"
                "\tTên đề tài: Thiết kế và xây dựng website bán cà phê tích hợp chatbot AI",
            )
        elif paragraph.text.startswith("Gíao viên hướng dẫn:"):
            replace_text(paragraph, "Giáo viên hướng dẫn: ThS. Vũ Thanh Nhàn")
        elif "Lớp học phần: 21DKTPM01" in paragraph.text:
            replace_text(
                paragraph,
                "Môn thi: Đồ án chuyên ngành kỹ thuật công nghệ thông tin"
                "                        Lớp học phần: 23DTH1A",
            )
        elif paragraph.text.startswith("Ngày thi:") and "Phòng thi:" in paragraph.text:
            replace_text(
                paragraph,
                "Ngày thi: ……/……/2026\tPhòng thi: ……………",
            )

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if "NĂM HỌC 2024 - 2025" in run.text:
                            run.text = run.text.replace(
                                "NĂM HỌC 2024 - 2025",
                                "NĂM HỌC 2025 - 2026",
                            )


def style_document(document: Document) -> None:
    for section in document.sections:
        section.left_margin = Cm(3.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)

    for style in document.styles:
        if getattr(style, "font", None) is not None:
            style.font.name = "Times New Roman"
            style_rpr = style._element.get_or_add_rPr()
            style_rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
        if style.name in BODY_STYLES:
            style.font.size = Pt(13)
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_after = Pt(0)

    for paragraph in document.paragraphs:
        if paragraph.style.name in BODY_STYLES:
            paragraph.paragraph_format.line_spacing = 1.5
            for run in paragraph.runs:
                run.font.size = Pt(13)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.15
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run._element.get_or_add_rPr().rFonts.set(
                            qn("w:eastAsia"), "Times New Roman"
                        )


def add_hanging_reference(writer: Writer, text: str) -> None:
    paragraph = writer.paragraph(text, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    paragraph.paragraph_format.left_indent = Cm(0.8)
    paragraph.paragraph_format.first_line_indent = Cm(-0.8)
    paragraph.paragraph_format.line_spacing = 1.5


def insert_toc_row(document: Document, anchor, text: str, style: str) -> None:
    template = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == style and "\t" in paragraph.text
    )
    paragraph = document.add_paragraph(style=style)
    if paragraph._p.pPr is not None:
        paragraph._p.remove(paragraph._p.pPr)
    paragraph._p.insert(0, deepcopy(template._p.pPr))
    label, page = text.rsplit("\t", 1)
    paragraph.add_run(label)
    paragraph.add_run("\t")
    paragraph.add_run(page)
    anchor._p.addprevious(paragraph._p)


def replace_static_toc_and_table_list(document: Document) -> None:
    toc_start = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "toc 1"
        and (
            paragraph.text.startswith("KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
            or paragraph.text.startswith("CHƯƠNG 5. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
        )
    )
    figure_list = find_paragraph(document, "DANH MỤC HÌNH")
    current = toc_start._p
    while current is not figure_list._p:
        following = current.getnext()
        current.getparent().remove(current)
        current = following

    toc_rows = [
        ("CHƯƠNG 5. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN\t44", "toc 1"),
        ("5.1. Kết quả đạt được\t44", "toc 2"),
        ("5.2. Hướng phát triển\t45", "toc 2"),
        ("TÀI LIỆU THAM KHẢO\t46", "toc 1"),
        ("PHỤ LỤC\t48", "toc 1"),
        ("PHỤ LỤC A. PHẠM VI ROUTE VÀ API CHÍNH\t48", "toc 2"),
        ("PHỤ LỤC B. CẤU HÌNH TRIỂN KHAI TỐI THIỂU\t48", "toc 2"),
        ("PHỤ LỤC C. KẾT QUẢ KIỂM THỬ VÀ PHẠM VI DỮ LIỆU\t49", "toc 2"),
    ]
    for text, style in toc_rows:
        insert_toc_row(document, figure_list, text, style)

    abbreviation_list = find_paragraph(document, "DANH MỤC CHỮ VIẾT TẮT")
    existing_rows = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "toc 1"
        and paragraph.text.startswith(("Bảng 5.1.", "Bảng PL."))
    ]
    for paragraph in existing_rows:
        remove_paragraph(paragraph)
    table_rows = [
        "Bảng 5.1. Tổng hợp kết quả chính của đồ án\t44",
        "Bảng PL.1. Route frontend và API tiêu biểu\t48",
        "Bảng PL.2. Biến môi trường production quan trọng\t48",
        "Bảng PL.3. Quality gates tại thời điểm hoàn thành\t49",
    ]
    for text in table_rows:
        insert_toc_row(document, abbreviation_list, text, "toc 1")


def write_chapter_five(document: Document) -> None:
    try:
        chapter = find_paragraph(document, "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
    except ValueError:
        chapter = find_paragraph(document, "CHƯƠNG 5. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
    references = find_paragraph(document, "TÀI LIỆU THAM KHẢO")
    appendix = find_paragraph(document, "PHỤ LỤC")

    clear_between(chapter, references)
    clear_between(references, appendix)
    sentinel = document.add_paragraph("__REPORT_END__", style="Normal")
    clear_between(appendix, sentinel)

    chapter.text = "CHƯƠNG 5. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN"
    chapter.style = "CAP 1"
    chapter.paragraph_format.page_break_before = True
    references.style = "CAP 1"
    references.paragraph_format.page_break_before = True
    appendix.style = "CAP 1"
    appendix.paragraph_format.page_break_before = True

    chapter_writer = Writer(document, references)
    chapter_writer.heading("5.1. Kết quả đạt được")
    chapter_writer.paragraph(
        "Đồ án đã xây dựng hoàn chỉnh website DẤU VỊ phục vụ giới thiệu và bán cà phê Việt Nam "
        "có truy xuất nguồn gốc. Hệ thống cung cấp trang chủ, bộ sưu tập sáu sản phẩm, tìm kiếm và "
        "lọc, trang chi tiết, hồ sơ lô, Coffee Advisor, hướng dẫn pha, giỏ hàng, checkout COD demo, "
        "đăng ký, đăng nhập và chatbot tư vấn. Giao diện được triển khai responsive, có trạng thái "
        "loading, empty, error, not-found và các thành phần hỗ trợ accessibility, SEO."
    )
    chapter_writer.paragraph(
        "Về kỹ thuật, frontend Next.js App Router sử dụng TypeScript và repository abstraction để "
        "chuyển đổi giữa mock data và HTTP API. Backend FastAPI áp dụng SQLAlchemy, Alembic và "
        "PostgreSQL/pgvector; giá đơn hàng được tính lại ở server, phiên đăng nhập dùng cookie "
        "HttpOnly và mật khẩu được băm Argon2id. Hệ thống đã được đóng gói Docker, triển khai frontend "
        "trên Vercel, backend và cơ sở dữ liệu trên VNPT Cloud, kết nối HTTPS qua DuckDNS và Caddy."
    )
    chapter_writer.paragraph(
        "Chatbot đã được hiện thực bằng LangGraph theo mô hình semantic router có kiểm soát. Groq chỉ "
        "chọn một route trong schema; LangGraph mới kích hoạt đúng Coffee Retrieval, Traceability, "
        "Brew Knowledge hoặc Commerce Policy Tool. Kết quả retrieval BM25 và vector được hợp nhất, "
        "grounding theo product/lot ID thật trước khi sinh phản hồi. Khi LLM hoặc vector không khả "
        "dụng, hệ thống có deterministic fallback và không tự tạo sản phẩm hay chứng nhận ngoài dữ liệu."
    )
    chapter_writer.table(
        "Bảng 5.1. Tổng hợp kết quả chính của đồ án",
        ["Nhóm", "Kết quả đạt được"],
        [
            ("Frontend", "Next.js App Router, đầy đủ route nghiệp vụ, responsive, cart persistence, SEO và accessibility cơ bản."),
            ("Backend", "FastAPI REST API cho catalog, lot, advisor, chatbot, auth và đơn COD; validation và transaction server-side."),
            ("Dữ liệu", "PostgreSQL, Alembic, pgvector; 6 sản phẩm, 6 lot demo, 8 documents và 24 knowledge chunks."),
            ("AI", "LangGraph semantic routing, 4 tool node, BM25 + vector + RRF, grounding, audit log và fallback."),
            ("Triển khai", "Vercel, Docker Compose, VNPT Cloud và Caddy."),
        ],
        widths_cm=[3.0, 11.8],
        font_size=11.5,
    )

    chapter_writer.heading("5.2. Hướng phát triển")
    chapter_writer.paragraph(
        "Phiên bản hiện tại đáp ứng mục tiêu của đồ án cơ sở và tạo nền tảng rõ ràng cho giai đoạn "
        "phát triển sản phẩm. Các hướng ưu tiên tiếp theo gồm:"
    )
    chapter_writer.bullet(
        "Dữ liệu và truy xuất: ",
        "kết nối dữ liệu thật từ nhà cung cấp, quy trình xác minh evidence, lịch sử thay đổi và cơ chế ký/xác nhận hồ sơ lô.",
    )
    chapter_writer.bullet(
        "Thương mại và quản trị: ",
        "xây dựng Admin UI/API với RBAC, tồn kho theo số lượng, trạng thái đơn, coupon, thanh toán và tích hợp đơn vị vận chuyển.",
    )
    chapter_writer.bullet(
        "Chatbot và RAG: ",
        "mở rộng Knowledge Base, CMS và tác vụ re-index nền; bổ sung lịch sử hội thoại, streaming, bộ đánh giá retrieval/grounding và theo dõi chi phí Groq.",
    )
    chapter_writer.bullet(
        "Bảo mật và vận hành: ",
        "bổ sung audit trail quản trị, secret manager, monitoring/alert, backup tự động, restore drill và kiểm thử bảo mật định kỳ.",
    )
    chapter_writer.bullet(
        "Chất lượng sản phẩm: ",
        "thực hiện usability test với người dùng mục tiêu, đo Core Web Vitals, kiểm thử tải API và tiếp tục hoàn thiện WCAG 2.2 AA.",
    )
    chapter_writer.paragraph(
        "Việc phát triển cần tiếp tục giữ nguyên nguyên tắc trung thực dữ liệu: nội dung demo phải được "
        "công bố rõ, chatbot chỉ dẫn nguồn đã kiểm tra và các claim môi trường chỉ được hiển thị khi có "
        "evidence phù hợp. Đây là điều kiện quan trọng để hệ thống có thể chuyển từ sản phẩm trình diễn "
        "sang nền tảng thương mại đáng tin cậy."
    )

    reference_writer = Writer(document, appendix)
    references_data = [
        '[1] Vercel, “Next.js Documentation – App Router,” https://nextjs.org/docs/app (truy cập ngày 12/08/2026).',
        '[2] Meta Open Source, “React Documentation,” https://react.dev/ (truy cập ngày 12/08/2026).',
        '[3] Microsoft, “TypeScript Handbook,” https://www.typescriptlang.org/docs/handbook/intro.html (truy cập ngày 12/08/2026).',
        '[4] Tailwind Labs, “Tailwind CSS Documentation,” https://tailwindcss.com/docs (truy cập ngày 12/08/2026).',
        '[5] S. Ramírez, “FastAPI Documentation,” https://fastapi.tiangolo.com/ (truy cập ngày 12/08/2026).',
        '[6] SQLAlchemy Authors, “SQLAlchemy 2.0 Documentation,” https://docs.sqlalchemy.org/en/20/ (truy cập ngày 12/08/2026).',
        '[7] PostgreSQL Global Development Group, “PostgreSQL 17 Documentation,” https://www.postgresql.org/docs/17/ (truy cập ngày 12/08/2026).',
        '[8] pgvector Contributors, “pgvector: Open-source vector similarity search for Postgres,” https://github.com/pgvector/pgvector (truy cập ngày 12/08/2026).',
        '[9] LangChain, “LangGraph StateGraph Reference,” https://reference.langchain.com/python/langgraph/graph (truy cập ngày 12/08/2026).',
        '[10] Groq, “Responses API and Structured Outputs,” https://console.groq.com/docs/responses-api (truy cập ngày 12/08/2026).',
        '[11] Docker, “Docker Compose Documentation,” https://docs.docker.com/compose/ (truy cập ngày 12/08/2026).',
        '[12] Caddy Authors, “reverse_proxy Directive,” https://caddyserver.com/docs/caddyfile/directives/reverse_proxy (truy cập ngày 12/08/2026).',
        '[13] OWASP Foundation, “Authentication Cheat Sheet,” https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html (truy cập ngày 12/08/2026).',
        '[14] W3C, “Web Content Accessibility Guidelines (WCAG) 2.2,” https://www.w3.org/TR/WCAG22/ (truy cập ngày 12/08/2026).',
    ]
    for item in references_data:
        add_hanging_reference(reference_writer, item)

    appendix_writer = Writer(document, sentinel)
    appendix_writer.heading("PHỤ LỤC A. PHẠM VI ROUTE VÀ API CHÍNH")
    appendix_writer.table(
        "Bảng PL.1. Route frontend và API tiêu biểu",
        ["Nhóm", "Địa chỉ", "Mục đích"],
        [
            ("Trang công khai", "/, /story, /brew-guide", "Giới thiệu thương hiệu, câu chuyện và hướng dẫn pha."),
            ("Sản phẩm", "/shop, /shop/[slug]", "Tìm kiếm, lọc, xem variant và thêm vào giỏ."),
            ("Truy xuất", "/traceability, /traceability/[lotCode]", "Tra mã lô và xem coffee passport demo."),
            ("Tư vấn", "/advisor", "Quiz sáu bước và top 3 recommendation rule-based."),
            ("Tài khoản", "/login, /register", "Đăng ký, đăng nhập và quản lý phiên HttpOnly."),
            ("Mua hàng", "/cart, /checkout", "Giỏ persist và tạo đơn COD demo."),
            ("Catalog API", "GET /api/v1/products, /lots", "Danh sách, chi tiết sản phẩm và lot."),
            ("AI API", "POST /api/v1/assistant/messages", "LangGraph semantic routing và grounded response."),
            ("Health", "/health/live, /ready, /rag", "Giám sát tiến trình, database, vector và workflow AI."),
        ],
        widths_cm=[2.5, 5.1, 7.2],
        font_size=11.2,
    )

    appendix_writer.heading("PHỤ LỤC B. CẤU HÌNH TRIỂN KHAI TỐI THIỂU")
    appendix_writer.table(
        "Bảng PL.2. Biến môi trường production quan trọng",
        ["Biến", "Phạm vi", "Ý nghĩa"],
        [
            ("DATABASE_URL", "Backend", "Kết nối PostgreSQL nội bộ."),
            ("SESSION_SECRET", "Backend", "Khóa bí mật cho định danh phiên và rate-limit hash."),
            ("GROQ_API_KEY", "Backend", "Khóa Groq; không đưa lên frontend hoặc Git."),
            ("AI_ENABLED", "Backend", "Bật semantic router và grounded generation."),
            ("VECTOR_SEARCH_ENABLED", "Backend", "Bật FastEmbed và pgvector retrieval."),
            ("NEXT_PUBLIC_API_BASE_URL", "Frontend", "Đường dẫn same-origin /backend-api trên Vercel."),
            ("API_BASE_URL", "Vercel server", "API gốc https://dauvi-api.duckdns.org/api/v1."),
        ],
        widths_cm=[4.1, 2.6, 8.1],
        font_size=11.2,
    )
    appendix_writer.paragraph(
        "Quy trình cập nhật production: sao lưu PostgreSQL khi có migration quan trọng; build image "
        "backend; chạy Docker Compose với --wait; xác nhận health=healthy; reload minute_caddy; kiểm "
        "tra /health/rag qua localhost, Docker network và HTTPS public trước khi hoàn tất triển khai."
    )

    appendix_writer.heading("PHỤ LỤC C. KẾT QUẢ KIỂM THỬ VÀ PHẠM VI DỮ LIỆU")
    appendix_writer.table(
        "Bảng PL.3. Quality gates tại thời điểm hoàn thành",
        ["Hạng mục", "Kết quả"],
        [
            ("Backend Ruff", "Đạt, không có lỗi lint."),
            ("Backend Pytest", "20/20 test đạt."),
            ("Frontend ESLint và typecheck", "Đạt."),
            ("Frontend Vitest", "24/24 unit test đạt."),
            ("Playwright", "13/13 E2E đạt trên production build."),
            ("Next.js production build", "Đạt, sinh 15 route pattern App Router."),
            ("Production RAG", "LangGraph ready; BM25 + pgvector; 24/24 chunks có embedding."),
        ],
        widths_cm=[5.0, 9.8],
        font_size=11.2,
    )
    appendix_writer.paragraph(
        "Disclosure bắt buộc: “Dữ liệu lô và đơn vị sản xuất đang được mô phỏng cho mục đích trình "
        "diễn đồ án.” Sáu hồ sơ lot hiện là Demo Data; hệ thống không tạo rating, review, chứng nhận, "
        "carbon footprint hoặc claim môi trường khi chưa có evidence tương ứng."
    )

    remove_paragraph(sentinel)
    replace_static_toc_and_table_list(document)
    fix_preface_layout(document)
    fix_template_placeholders(document)
    ensure_demo_link(document)
    style_document(document)
    standardize_table_colors(document)
    set_update_fields(document)


def replace_diagrams(docx_path: Path, diagram_dir: Path) -> None:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with ZipFile(docx_path, "r") as source, ZipFile(
            temporary, "w", compression=ZIP_DEFLATED
        ) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                replacement = DIAGRAM_MEDIA.get(item.filename)
                if replacement:
                    image_path = diagram_dir / replacement
                    with Image.open(image_path) as image:
                        image = image.convert("RGB")
                        image.thumbnail((2000, 1125), Image.Resampling.LANCZOS)
                        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as png_handle:
                            png_path = Path(png_handle.name)
                        try:
                            image.save(png_path, format="PNG", optimize=True)
                            data = png_path.read_bytes()
                        finally:
                            png_path.unlink(missing_ok=True)
                target.writestr(item, data)
        os.replace(temporary, docx_path)
    finally:
        temporary.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("target", type=Path)
    parser.add_argument("--final", type=Path)
    args = parser.parse_args()

    target = args.target.resolve()
    final = (args.final or target.with_name(f"{target.stem}_FINAL.docx")).resolve()
    backup_dir = target.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f"{target.stem}.before-ch5-{datetime.now():%Y%m%d-%H%M%S}.docx"
    shutil.copy2(target, backup)

    document = Document(target)
    write_chapter_five(document)
    with tempfile.NamedTemporaryFile(dir=target.parent, suffix=".tmp.docx", delete=False) as handle:
        temporary = Path(handle.name)
    document.save(temporary)
    os.replace(temporary, target)

    diagram_dir = target.parent / "report-assets" / "diagrams"
    replace_diagrams(target, diagram_dir)
    shutil.copy2(target, final)
    print(target)
    print(final)
    print(backup)


if __name__ == "__main__":
    main()
